"""Teacher slide decks (one English deck, one Spanish deck per unit).

Accessibility choices: every slide uses a layout title placeholder with a unique title, body text in the
body placeholder (correct reading order), 24pt+ body, dark text on white, language tagged on every run,
teacher notes carry timing and facilitation notes. No images, so no alt text is needed.
"""
import re
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn

from content import COURSE, DAY_SHAPE, LAB_SHAPE

GREEN = RGBColor(0x1B, 0x5E, 0x3A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
LABEL = {
    "en": {"unit": "Unit", "day": "Day", "lab": "Lab day", "obj": "Learning targets", "vocab": "Vocabulary",
           "bell": "Bellringer", "inst": "Direct instruction", "guided": "Guided practice", "indep": "Independent practice",
           "labw": "Lab", "exit": "Exit ticket", "recipe": "Recipe", "plan": "The 48 minutes", "min": "min",
           "roles": "Group jobs", "of": "of", "lang": "en-US"},
    "es": {"unit": "Unidad", "day": "Día", "lab": "Día de laboratorio", "obj": "Metas de aprendizaje", "vocab": "Vocabulario",
           "bell": "Actividad de inicio", "inst": "Instrucción directa", "guided": "Práctica guiada",
           "indep": "Práctica independiente", "labw": "Laboratorio", "exit": "Boleto de salida", "recipe": "Receta",
           "plan": "Los 48 minutos", "min": "min", "roles": "Tareas del grupo", "of": "de", "lang": "es-US"},
}
SPLIT = re.compile(r"(?<![0-9]\.)(?<=[.!?])\s+(?=[A-ZÁÉÍÓÚ¿¡0-9])")


def _chunks(text, limit=230):
    sents = SPLIT.split(text)
    out, cur = [], []
    n = 0
    for s in sents:
        if cur and n + len(s) > limit:
            out.append(cur)
            cur, n = [], 0
        cur.append(s)
        n += len(s)
    if cur:
        out.append(cur)
    return out


class Deck:
    def __init__(self, title, lang):
        self.p = Presentation()
        self.p.slide_width = Inches(13.333)
        self.p.slide_height = Inches(7.5)
        self.lang = lang
        self.L = LABEL[lang]
        self.p.core_properties.title = title
        self.p.core_properties.language = self.L["lang"]
        self.p.core_properties.author = f"{COURSE['school']} {COURSE['department']}"
        self.titles = set()

    def _tag(self, tf):
        for para in tf.paragraphs:
            for r in para.runs:
                r._r.get_or_add_rPr().set("lang", self.L["lang"])

    def slide(self, title, bullets, notes="", size=26, numbered=False):
        base = title
        k = 2
        while title in self.titles:
            title = f"{base} ({k})"
            k += 1
        self.titles.add(title)
        s = self.p.slides.add_slide(self.p.slide_layouts[1])
        t, body = s.shapes.title, s.placeholders[1]
        t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.2)
        body.left, body.top, body.width, body.height = Inches(0.8), Inches(1.7), Inches(11.7), Inches(5.3)
        t.text_frame.text = title
        for r in t.text_frame.paragraphs[0].runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(32), True, GREEN, "Arial"
        tf = body.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = (f"{i + 1}. " if numbered else "") + b
            para.space_after = Pt(10)
            for r in para.runs:
                r.font.size, r.font.color.rgb, r.font.name = Pt(size), INK, "Arial"
        self._tag(t.text_frame)
        self._tag(tf)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def text_slides(self, title, text, notes="", size=26):
        parts = _chunks(text)
        for i, c in enumerate(parts):
            t = title if len(parts) == 1 else f"{title} ({i + 1} {self.L['of']} {len(parts)})"
            self.slide(t, c, notes, size)


def build_deck(unit, lang, path):
    L = LABEL[lang]
    g = lambda x: x[lang]
    d = Deck(f"{g(COURSE['title'])} - {L['unit']} {unit['number']}: {g(unit['title'])}", lang)
    d.slide(f"{L['unit']} {unit['number']}: {g(unit['title'])}", [g(unit["essential_question"]), g(unit["overview"])],
            "Opening slide. Read the essential question aloud and connect it to the certification objectives: " + ", ".join(unit["cert"]))
    vl = {v["term"]["en"]: v for v in unit["vocab_bank"]}
    for les in unit["lessons"]:
        is_lab = les["kind"] == "lab"
        pre = f"{L['lab'] if is_lab else L['day']} {les['day']}"
        tm = les["teacher"]
        notes_all = "Materials: " + "; ".join(tm["materials"]) + "\n" + "\n".join(tm["notes"])
        shape = LAB_SHAPE if is_lab else DAY_SHAPE
        d.slide(f"{pre}: {g(les['title'])}",
                [f"{m} {L['min']}: {es if lang == 'es' else en}" for m, en, es in shape], notes_all, size=26)
        d.slide(f"{pre} - {L['obj']}", [g(o) for o in les["objectives"]], "Read the targets aloud. Students restate one in their own words.")
        terms = [vl[t] for t in les.get("vocab", []) if t in vl]
        if terms:
            d.slide(f"{pre} - {L['vocab']}", [f"{g(v['term'])}: {g(v['definition'])}" for v in terms], "Preview vocabulary; ask for examples.")
        d.slide(f"{pre} - {L['bell']}", [g(les["bellringer"])], "5 minutes. Collect on the exit ticket sheet or discuss as a class.")
        for blk in les["instruction"]:
            d.text_slides(f"{pre} - {g(blk['heading'])}", g(blk["body"]), "Check for understanding after this slide.")
        gd = les["guided"]
        d.slide(f"{pre} - {L['guided']}: {g(gd['title'])}", [g(s) for s in gd["steps"]], "Model the first step, then release.", numbered=True)
        ind = les["independent"]
        d.slide(f"{pre} - {L['labw'] if ind['kind'] == 'lab' else L['indep']}: {g(ind['title'])}",
                [g(s) for s in ind["steps"]], "Circulate; Safety Officers check the checkpoints.", numbered=True)
        d.slide(f"{pre} - {L['exit']}", [g(x["q"]) for x in les["exit"]],
                "Answer key: " + " | ".join(x["answer"]["en"] for x in les["exit"]), numbered=True)
    for r in unit["recipes"].values():
        d.slide(f"{L['recipe']}: {g(r['title'])}",
                [f"{i['qty']} {g(i['unit'])} {g(i['item'])}" for i in r["ingredients"]],
                "Ingredients display for lab setup. Peanut-free kitchen: confirm labels.")
    d.p.save(path)
