"""Automated accessibility checks for the generated site, Word files, and slide decks.

Automated checks catch structure problems; they do not replace a manual pass with a screen reader
(NVDA/VoiceOver) and Microsoft's Accessibility Checker. The report says what was and was not tested.
"""
import glob
import os
import re
from html.parser import HTMLParser

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Pt


def lum(hexc):
    hexc = hexc.lstrip("#")
    r, g, b = [int(hexc[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    f = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def contrast(a, b):
    la, lb = sorted((lum(a), lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


PAIRS_LIGHT = [("text on background", "#1a1a1a", "#ffffff"), ("muted text", "#4a4a4a", "#ffffff"),
               ("header text", "#ffffff", "#1b5e3a"), ("link", "#0b4f8a", "#ffffff"),
               ("link on panel", "#0b4f8a", "#f2f7f3"), ("note text", "#5c3200", "#fff4e0"),
               ("heading green on white", "#1b5e3a", "#ffffff")]
PAIRS_DARK = [("text on background", "#f2f2f2", "#121512"), ("muted text", "#cfd6d1", "#121512"),
              ("header text", "#0c1f14", "#8fd3a7"), ("link", "#9cc9f5", "#121512"),
              ("link on panel", "#9cc9f5", "#1c241f"), ("note text", "#ffd9a3", "#2c2111")]


class Page(HTMLParser):
    def __init__(self):
        super().__init__()
        self.h, self.tables, self.problems = [], [], []
        self.lang = None
        self.title = ""
        self.in_title = False
        self.skip = False
        self.cur = None
        self.imgs_no_alt = 0
        self.buttons = 0
        self.bad_lang = 0

    def handle_starttag(self, tag, a):
        a = dict(a)
        if tag == "html":
            self.lang = a.get("lang")
        if tag == "title":
            self.in_title = True
        if re.fullmatch(r"h[1-6]", tag):
            self.h.append(int(tag[1]))
        if tag == "a" and a.get("class") == "skip":
            self.skip = True
        if tag == "table":
            self.cur = {"caption": False, "th_scope": 0, "th": 0}
            self.tables.append(self.cur)
        if tag == "caption" and self.cur is not None:
            self.cur["caption"] = True
        if tag == "th" and self.cur is not None:
            self.cur["th"] += 1
            if a.get("scope"):
                self.cur["th_scope"] += 1
        if tag == "img" and not a.get("alt"):
            self.imgs_no_alt += 1
        if tag == "span" and a.get("lang") not in (None, "en", "es"):
            self.bad_lang += 1

    def handle_endtag(self, tag):
        if tag == "title":
            self.in_title = False
        if tag == "table":
            self.cur = None

    def handle_data(self, d):
        if self.in_title:
            self.title += d


def audit_html(site_dir):
    res = []
    for f in sorted(glob.glob(os.path.join(site_dir, "*.html"))):
        p = Page()
        p.feed(open(f, encoding="utf-8").read())
        name = os.path.basename(f)
        checks = [
            ("html lang set", bool(p.lang)), ("page title present", bool(p.title.strip())), ("skip link present", p.skip),
            ("exactly one h1", p.h.count(1) == 1),
            ("no skipped heading levels", all(b - a <= 1 for a, b in zip(p.h, p.h[1:]))),
            ("every table has caption", all(t["caption"] for t in p.tables)),
            ("every table has header cells with scope", all(t["th"] and t["th"] == t["th_scope"] for t in p.tables)),
            ("images have alt text", p.imgs_no_alt == 0),
        ]
        res.append((name, checks))
    return res


def audit_docx(path):
    d = Document(path)
    checks = []
    checks.append(("document title set", bool(d.core_properties.title)))
    checks.append(("document language set", bool(d.core_properties.language)))
    heads = [p for p in d.paragraphs if p.style.name.startswith("Heading")]
    checks.append(("uses heading styles", len(heads) > 0))
    lv = [int(p.style.name.split()[-1]) for p in heads]
    checks.append(("no skipped heading levels", all(b - a <= 1 for a, b in zip(lv, lv[1:]))))
    empties = [p for p in d.paragraphs if not p.text.strip() and not p._p.xpath(".//w:br")]
    checks.append(("no empty paragraphs", len(empties) == 0))
    ok_t = all(t._tbl.xpath("./w:tr[1]/w:trPr/w:tblHeader") and t._tbl.tblPr.find(qn("w:tblCaption")) is not None
               and t._tbl.tblPr.find(qn("w:tblDescription")) is not None for t in d.tables)
    checks.append(("every table has header row, caption, description", ok_t))
    imgs = d.element.body.xpath(".//wp:docPr")
    checks.append(("images have alt text", all((i.get("descr") or "").strip() for i in imgs)))
    sizes = [r.font.size.pt for p in d.paragraphs for r in p.runs if r.font.size]
    checks.append(("no text under 11pt", all(s >= 11 for s in sizes)))
    return checks


def audit_pptx(path):
    pr = Presentation(path)
    checks, titles = [], []
    no_title = no_notes = small = 0
    for s in pr.slides:
        t = s.shapes.title
        if t is None or not t.text_frame.text.strip():
            no_title += 1
        else:
            titles.append(t.text_frame.text)
        if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()):
            no_notes += 1
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size < Pt(18):
                            small += 1
    checks.append(("every slide has a title", no_title == 0))
    checks.append(("slide titles are unique", len(titles) == len(set(titles))))
    checks.append(("every slide has speaker notes", no_notes == 0))
    checks.append(("no text under 18pt", small == 0))
    checks.append(("presentation title and language set", bool(pr.core_properties.title) and bool(pr.core_properties.language)))
    return checks


def run(site_dir, dl_dir, report_path):
    lines = ["# Accessibility audit (automated)", "",
             "Standard targeted: WCAG 2.1 AA / Section 508. This report covers automated structural checks only.", ""]
    fails = 0
    lines += ["## Color contrast (WCAG AA needs 4.5:1 for normal text)", "", "| Pair | Mode | Ratio | Result |", "|---|---|---|---|"]
    for mode, pairs in (("light", PAIRS_LIGHT), ("dark", PAIRS_DARK)):
        for nm, fg, bg in pairs:
            r = contrast(fg, bg)
            ok = r >= 4.5
            fails += not ok
            lines.append(f"| {nm} ({fg} on {bg}) | {mode} | {r:.1f}:1 | {'PASS' if ok else 'FAIL'} |")
    lines += ["", "## Website", ""]
    for name, checks in audit_html(site_dir):
        for c, ok in checks:
            fails += not ok
        bad = [c for c, ok in checks if not ok]
        lines.append(f"- `{name}`: {'all ' + str(len(checks)) + ' checks passed' if not bad else 'FAILED: ' + ', '.join(bad)}")
    lines += ["", "## Word documents", ""]
    for f in sorted(glob.glob(os.path.join(dl_dir, "*.docx"))):
        checks = audit_docx(f)
        bad = [c for c, ok in checks if not ok]
        fails += len(bad)
        lines.append(f"- `{os.path.basename(f)}`: {'all ' + str(len(checks)) + ' checks passed' if not bad else 'FAILED: ' + ', '.join(bad)}")
    lines += ["", "## Slide decks", ""]
    for f in sorted(glob.glob(os.path.join(dl_dir, "*.pptx"))):
        checks = audit_pptx(f)
        bad = [c for c, ok in checks if not ok]
        fails += len(bad)
        lines.append(f"- `{os.path.basename(f)}`: {'all ' + str(len(checks)) + ' checks passed' if not bad else 'FAILED: ' + ', '.join(bad)}")
    lines += ["", "## Not covered by this automated audit (do before distributing)", "",
              "- Run Microsoft Office Accessibility Checker on each .docx and .pptx.",
              "- Read one page of the site and one Word file with a screen reader (NVDA on Windows, VoiceOver on Mac/iOS) in both languages.",
              "- Export PDFs from Word with the 'Document structure tags for accessibility' option checked, then run Adobe's PDF accessibility check. This build does not generate PDFs.",
              "- Add captions and a transcript for any video or audio you add later.",
              "- Spanish text needs review by a fluent speaker, including the `lang` tags on every passage.",
              "- Keyboard-only walkthrough of the site (Tab through the language buttons and links).", "",
              f"**Result:** {'PASS - no automated failures' if not fails else str(fails) + ' automated failure(s) - see above'}", ""]
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return fails
