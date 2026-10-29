"""Accessible Word documents: student workbook (one language per file), recipe cards, teacher guide.

Accessibility choices: built-in Heading styles for structure, real list styles, header rows marked on every
table with caption/description, document language and title set, 12pt+ Arial, no empty paragraphs for spacing,
answer spaces are labeled table cells (not underscores), no color-only meaning.
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

from content import COURSE, DAY_SHAPE, LAB_SHAPE, ROLES, SUPPORTS, CERT_OBJECTIVES

LABELS = {
    "en": {
        "name": "Name", "date": "Date", "period": "Class period", "day": "Day", "lab_day": "Lab day",
        "objectives": "Learning targets", "vocab": "Vocabulary", "term": "Term", "definition": "Definition",
        "bellringer": "Bellringer", "instruction": "Direct instruction", "guided": "Guided practice",
        "independent": "Independent practice", "lab": "Lab", "exit": "Exit ticket", "reading": "Reading",
        "reading_q": "Reading check", "answer": "Your answer", "recipe": "Recipe card", "serves": "Serves",
        "minutes": "Active time (minutes)", "ingredients": "Ingredients", "equipment": "Equipment", "steps": "Steps",
        "checkpoint": "Safety checkpoint", "allergens": "Allergy note", "cost": "Estimated cost per group",
        "qty": "Amount", "unit": "Measure", "item": "Ingredient", "roles": "Group jobs", "role": "Job",
        "does": "What you do", "essential": "Essential question", "overview": "Unit overview",
        "plan": "The 48 minutes", "min": "Minutes", "segment": "Segment", "unit_word": "Unit",
        "station": "Station check", "check": "Done?", "item_c": "Check", "initials": "Initials",
        "yes": "Yes / No", "work_plan": "Group work plan", "who": "Who", "task": "Task", "time": "Minutes",
        "scaled": "Your group's amount", "lang": "en-US", "review": "Spanish text is machine-assisted and awaiting fluent review.",
    },
    "es": {
        "name": "Nombre", "date": "Fecha", "period": "Período de clase", "day": "Día", "lab_day": "Día de laboratorio",
        "objectives": "Metas de aprendizaje", "vocab": "Vocabulario", "term": "Término", "definition": "Definición",
        "bellringer": "Actividad de inicio", "instruction": "Instrucción directa", "guided": "Práctica guiada",
        "independent": "Práctica independiente", "lab": "Laboratorio", "exit": "Boleto de salida", "reading": "Lectura",
        "reading_q": "Comprueba tu lectura", "answer": "Tu respuesta", "recipe": "Tarjeta de receta", "serves": "Rinde",
        "minutes": "Tiempo activo (minutos)", "ingredients": "Ingredientes", "equipment": "Equipo", "steps": "Pasos",
        "checkpoint": "Punto de control de seguridad", "allergens": "Nota sobre alergias",
        "cost": "Costo estimado por grupo", "qty": "Cantidad", "unit": "Medida", "item": "Ingrediente",
        "roles": "Tareas del grupo", "role": "Tarea", "does": "Qué haces", "essential": "Pregunta esencial",
        "overview": "Resumen de la unidad", "plan": "Los 48 minutos", "min": "Minutos", "segment": "Segmento",
        "unit_word": "Unidad", "station": "Revisión de la estación", "check": "¿Listo?", "item_c": "Revisar",
        "initials": "Iniciales", "yes": "Sí / No", "work_plan": "Plan de trabajo del grupo", "who": "Quién",
        "task": "Tarea", "time": "Minutos", "scaled": "La cantidad de tu grupo", "lang": "es-US",
        "review": "El texto en español fue elaborado con ayuda de traducción automática y está pendiente de revisión por un hablante fluido.",
    },
}

HEAD = RGBColor(0x1B, 0x5E, 0x3A)  # dark green, 7.6:1 on white
BODY = RGBColor(0x1A, 0x1A, 0x1A)


def _lang_xml(el, tag, lang):
    pr = el.find(qn(tag))
    if pr is None:
        pr = OxmlElement(tag)
        el.append(pr)
    for old in pr.findall(qn("w:lang")):
        pr.remove(old)
    l = OxmlElement("w:lang")
    l.set(qn("w:val"), lang)
    l.set(qn("w:eastAsia"), lang)
    l.set(qn("w:bidi"), "ar-SA")
    pr.append(l)


def new_doc(title, lang_code, subject=""):
    d = Document()
    cp = d.core_properties
    cp.title = title
    cp.subject = subject
    cp.author = f"{COURSE['school']} {COURSE['department']}"
    cp.language = lang_code
    cp.keywords = "Whole Foods and Nutrition; accessible"
    st = d.styles
    base = st["Normal"]
    base.font.name = "Arial"
    base.font.size = Pt(12)
    base.font.color.rgb = BODY
    base.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    base.paragraph_format.space_after = Pt(6)
    base.paragraph_format.line_spacing = 1.25
    # document default language
    rpr_def = st.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault"))
    _lang_xml(rpr_def, "w:rPr", lang_code)
    for name, size in (("Title", 26), ("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 13)):
        s = st[name]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = HEAD
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        s.element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        s.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    for nm in ("List Number", "List Bullet", "Heading 1", "Heading 2", "Heading 3"):
        st[nm].paragraph_format.space_before = Pt(10 if nm.startswith("List") else 14)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    return d


def _set_cell_shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def table(d, headers, rows, caption, description, widths=None, row_height=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    cap = OxmlElement("w:tblCaption")
    cap.set(qn("w:val"), caption)
    tblPr.append(cap)
    desc = OxmlElement("w:tblDescription")
    desc.set(qn("w:val"), description)
    tblPr.append(desc)
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        _set_cell_shade(c, "E3EFE7")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    if row_height:
        for row in t.rows[1:]:
            row.height = Inches(row_height)
    return t


def answer_box(d, L, lines=3):
    table(d, [L["answer"]], [[""]], L["answer"], "Space for the student to write an answer.", widths=[6.5],
          row_height=0.35 * lines)


def bullets(d, items, style="List Bullet"):
    for it in items:
        d.add_paragraph(it, style=style)


def _t(x, lang):
    return x[lang]


def _recipe(d, r, L, lang):
    d.add_heading(f"{L['recipe']}: {_t(r['title'], lang)}", level=2)
    d.add_paragraph(f"{L['serves']}: {r['serves']}   |   {L['minutes']}: {r['active_minutes']}")
    p = d.add_paragraph()
    p.add_run(f"{L['allergens']}: ").bold = True
    p.add_run(_t(r["allergens"], lang))
    if r.get("scale_note"):
        p = d.add_paragraph()
        p.add_run(_t(r["scale_note"], lang))
    if r.get("scale_col"):
        table(d, [L["qty"], L["unit"], L["item"], L["scaled"]],
              [[i["qty"], _t(i["unit"], lang), _t(i["item"], lang), ""] for i in r["ingredients"]],
              L["ingredients"], "Ingredients with amount, measure and name, plus a blank column for the student's scaled amount.",
              widths=[0.8, 1.1, 2.6, 2.0], row_height=0.35)
    else:
        table(d, [L["qty"], L["unit"], L["item"]],
              [[i["qty"], _t(i["unit"], lang), _t(i["item"], lang)] for i in r["ingredients"]],
              L["ingredients"], "Ingredients with amount, measure and name.", widths=[1.0, 1.4, 4.1])
    d.add_heading(L["equipment"], level=3)
    bullets(d, [_t(x, lang) for x in r["equipment"]])
    d.add_heading(L["steps"], level=3)
    for s in r["steps"]:
        d.add_paragraph(_t(s["text"], lang), style="List Number")
        if s["checkpoint"]:
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.add_run(f"{L['checkpoint']}: ").bold = True
            p.add_run(_t(s["checkpoint"], lang))
    p = d.add_paragraph()
    p.add_run(f"{L['cost']}: ").bold = True
    p.add_run(r["cost_per_group"])


def build_workbook(unit, lang, path):
    L = LABELS[lang]
    title = f"{_t(COURSE['title'], lang)} - {L['unit_word']} {unit['number']}: {_t(unit['title'], lang)}"
    d = new_doc(title, L["lang"], "Student workbook")
    d.add_paragraph(title, style="Title")
    table(d, [L["name"], L["date"], L["period"]], [["", "", ""]], L["name"], "Student name, date, and class period.",
          widths=[3.0, 1.75, 1.75], row_height=0.4)
    d.add_heading(L["overview"], level=1)
    p = d.add_paragraph()
    p.add_run(f"{L['essential']}: ").bold = True
    p.add_run(_t(unit["essential_question"], lang))
    d.add_paragraph(_t(unit["overview"], lang))
    table(d, [L["role"], L["does"]], [[_t(r["role"], lang), _t(r["job"], lang)] for r in ROLES],
          L["roles"], "The four group jobs and what each person does.", widths=[2.2, 4.3])
    d.add_heading(L["vocab"], level=1)
    table(d, [L["term"], L["definition"]], [[_t(v["term"], lang), _t(v["definition"], lang)] for v in unit["vocab_bank"]],
          L["vocab"], "Unit vocabulary terms and definitions.", widths=[2.2, 4.3])
    rd = unit["reading"]
    d.add_heading(f"{L['reading']}: {_t(rd['title'], lang)}", level=1)
    for para in rd["body"]:
        d.add_paragraph(_t(para, lang))
    d.add_heading(L["reading_q"], level=2)
    for q in rd["questions"]:
        d.add_paragraph(_t(q["q"], lang), style="List Number")
        answer_box(d, L, 2)
    vl = {v["term"]["en"]: v for v in unit["vocab_bank"]}
    for les in unit["lessons"]:
        is_lab = les["kind"] == "lab"
        d.add_heading(f"{L['lab_day'] if is_lab else L['day']} {les['day']}: {_t(les['title'], lang)}", level=1)
        shape = LAB_SHAPE if is_lab else DAY_SHAPE
        table(d, [L["min"], L["segment"]], [[m, es if lang == "es" else en] for m, en, es in shape],
              L["plan"], "Minutes for each part of the class.", widths=[1.0, 5.5])
        d.add_heading(L["objectives"], level=2)
        bullets(d, [_t(o, lang) for o in les["objectives"]])
        d.add_heading(L["bellringer"], level=2)
        d.add_paragraph(_t(les["bellringer"], lang))
        answer_box(d, L, 3)
        d.add_heading(L["instruction"], level=2)
        for blk in les["instruction"]:
            d.add_heading(_t(blk["heading"], lang), level=3)
            d.add_paragraph(_t(blk["body"], lang))
        g = les["guided"]
        d.add_heading(f"{L['guided']}: {_t(g['title'], lang)}", level=2)
        for s in g["steps"]:
            d.add_paragraph(_t(s, lang), style="List Number")
        answer_box(d, L, 4)
        ind = les["independent"]
        d.add_heading(f"{L['lab'] if ind['kind'] == 'lab' else L['independent']}: {_t(ind['title'], lang)}", level=2)
        for s in ind["steps"]:
            d.add_paragraph(_t(s, lang), style="List Number")
        if ind["kind"] != "lab":
            answer_box(d, L, 5)
        else:
            d.add_heading(L["work_plan"], level=3)
            table(d, [L["who"], L["task"], L["time"]], [["", "", ""]] * 5, L["work_plan"],
                  "Table for the group to plan who does each task and how long it takes.", widths=[1.8, 3.7, 1.0], row_height=0.4)
        d.add_heading(L["exit"], level=2)
        for x in les["exit"]:
            d.add_paragraph(_t(x["q"], lang), style="List Number")
            answer_box(d, L, 2)
    for r in unit["recipes"].values():
        d.add_page_break()
        _recipe(d, r, L, lang)
    if lang == "es":
        d.add_paragraph(L["review"])
    d.save(path)


def build_recipe_cards(unit, lang, path):
    L = LABELS[lang]
    title = f"{_t(COURSE['title'], lang)} - {L['recipe']} ({L['unit_word']} {unit['number']})"
    d = new_doc(title, L["lang"], "Recipe cards")
    d.add_paragraph(title, style="Title")
    first = True
    for r in unit["recipes"].values():
        if not first:
            d.add_page_break()
        first = False
        _recipe(d, r, L, lang)
        d.add_heading(L["station"], level=3)
        rows = [[_t(x, lang), "", ""] for x in (
            {"en": "Hands washed 20 seconds", "es": "Manos lavadas 20 segundos"},
            {"en": "Counter washed, rinsed, sanitized", "es": "Mostrador lavado, enjuagado, desinfectado"},
            {"en": "Raw and ready-to-eat foods kept apart", "es": "Alimentos crudos y listos para comer separados"},
            {"en": "Cold foods back in refrigerator", "es": "Alimentos fríos de vuelta en el refrigerador"},
            {"en": "Dishes washed and put away", "es": "Platos lavados y guardados"},
            {"en": "Equipment returned", "es": "Equipo devuelto"})]
        table(d, [L["item_c"], L["yes"], L["initials"]], rows, L["station"],
              "Checklist the group completes before cleanup is approved.", widths=[4.0, 1.2, 1.3])
    d.save(path)


def build_teacher_guide(unit, path):
    L = LABELS["en"]
    title = f"Teacher Guide - Unit {unit['number']}: {unit['title']['en']}"
    d = new_doc(title, "en-US", "Teacher guide")
    d.add_paragraph(title, style="Title")
    p = d.add_paragraph()
    p.add_run("Days: ").bold = True
    p.add_run(f"{unit['days']}   ")
    p.add_run("Labs: ").bold = True
    p.add_run(str(unit["labs"]))
    d.add_heading("Certification alignment", level=1)
    bullets(d, [f"{k}: {CERT_OBJECTIVES[k]}" for k in unit["cert"]])
    d.add_heading("Unit overview", level=1)
    d.add_paragraph(unit["overview"]["en"])
    d.add_heading("Day-by-day", level=1)
    for les in unit["lessons"]:
        d.add_heading(f"Day {les['day']}: {les['title']['en']}", level=2)
        tm = les["teacher"]
        p = d.add_paragraph()
        p.add_run("Cert objectives: ").bold = True
        p.add_run(", ".join(les["cert"]))
        d.add_heading("Materials", level=3)
        bullets(d, tm["materials"])
        d.add_heading("Teaching notes", level=3)
        bullets(d, tm["notes"])
        d.add_heading("Exit ticket answer key", level=3)
        for x in les["exit"]:
            d.add_paragraph(f"{x['q']['en']}  Answer: {x['answer']['en']}", style="List Number")
    d.add_heading("Reading check answer key", level=1)
    for q in unit["reading"]["questions"]:
        d.add_paragraph(f"{q['q']['en']}  Answer: {q['answer']['en']}", style="List Number")
    d.add_heading("Unit quiz answer key", level=1)
    letters = "ABCD"
    rows = [[i + 1, q["cert"], letters[q["answer"]] + ". " + q["choices"][q["answer"]]["en"]] for i, q in enumerate(unit["quiz"])]
    table(d, ["#", "Cert objective", "Correct answer"], rows, "Quiz answer key", "Question number, certification objective and correct answer.",
          widths=[0.5, 1.2, 4.8])
    d.add_heading("Practical rubric", level=1)
    pr = unit["practical"]
    d.add_paragraph(pr["scale"])
    table(d, ["Skill", "Strong (8-10)", "Developing (4-7)", "Beginning (0-3)"],
          [[c["name"]["en"], c["high"]["en"], c["mid"]["en"], c["low"]["en"]] for c in pr["criteria"]],
          "Unit practical rubric", "Four skills scored on three performance levels.", widths=[1.5, 1.7, 1.7, 1.6])
    d.add_heading("Supports and accessibility", level=1)
    bullets(d, SUPPORTS)
    d.add_heading("Before you use this unit", level=1)
    bullets(d, [
        "Spanish text is machine-assisted and has not been reviewed by a fluent speaker. Have one review before distributing.",
        "Cost and nutrition estimates are approximate. Check current prices and labels.",
        "Food safety temperatures follow USDA; confirm your school's sanitizer and health-code requirements.",
        "Course policies on Day 1 come from the current syllabus; update them here if the syllabus changes.",
    ])
    d.save(path)
