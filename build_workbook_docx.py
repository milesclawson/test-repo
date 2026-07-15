"""
Shared student workbook (.docx) builder.
Given a course config + units (unit meta + lessons), emits a printable
student workbook: cover, how-to-use, then per unit a divider page + vocab
bank, then one workbook page per lesson (objectives, vocab, warm-up, guided
notes cues, activity, reading + questions, quiz with answer space, exit
ticket), then a unit assessment page.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def para(doc, text="", size=11, bold=False, italic=False, color=None,
         align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return p


def heading_bar(doc, text, hex_color, text_color=RGBColor(0xFF, 0xFF, 0xFF), size=13):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = True
    cell = t.rows[0].cells[0]
    set_cell_bg(cell, hex_color)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = text_color
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def blank_lines(doc, n=3, color=RGBColor(0x99, 0x99, 0x99)):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "999999")
        pbdr.append(bottom)
        # w:pBdr must precede w:spacing/w:ind/w:jc in CT_PPrBase's required sequence
        pPr.insert(0, pbdr)


def build_workbook(*, output_path, course_title, school_name, grade, semester_label,
                    primary_hex, primary_dk_hex, accent_hex, units, intro_text):
    doc = Document()
    zoom_el = doc.settings.element.find(qn("w:zoom"))
    if zoom_el is not None and zoom_el.get(qn("w:percent")) is None:
        zoom_el.set(qn("w:percent"), "100")
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ---- Cover ----
    para(doc, school_name.upper(), size=12, bold=True, color=RGBColor.from_string(primary_dk_hex.lstrip("#")),
         align=WD_ALIGN_PARAGRAPH.CENTER, space_before=140)
    para(doc, course_title, size=32, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    para(doc, "Student Workbook", size=18, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         color=RGBColor.from_string(primary_hex.lstrip("#")), space_after=20)
    para(doc, f"{grade} · {semester_label}", size=13, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x66, 0x66, 0x66))
    para(doc, "\n\nName: _____________________________________          Period: __________",
         size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=60)
    doc.add_page_break()

    # ---- How to use ----
    heading_bar(doc, "HOW TO USE THIS WORKBOOK", primary_hex.lstrip("#") if False else primary_hex.replace("#", ""))
    para(doc, intro_text, size=11.5)
    para(doc, "Each lesson page gives you: today's vocabulary, a warm-up, guided note-taking "
              "cues, the class activity, a reading when one is assigned, a short quiz, and an "
              "exit ticket. Bring this workbook to class every day.", size=11.5)
    doc.add_page_break()

    primary_hexc = primary_hex.replace("#", "")
    primary_dk_hexc = primary_dk_hex.replace("#", "")
    accent_hexc = accent_hex.replace("#", "")

    for u in units:
        meta, lessons = u["meta"], u["lessons"]
        # Unit divider
        heading_bar(doc, f"UNIT {meta['number']} · {meta['name'].upper()}", primary_dk_hexc, size=16)
        para(doc, meta.get("essential_question", ""), italic=True, size=13,
             color=RGBColor.from_string(primary_hexc))
        para(doc, meta.get("overview", ""), size=11.5)
        para(doc, f"Days {meta['days'][0]}–{meta['days'][1]}  ·  {len(lessons)} lessons", size=10.5,
             color=RGBColor(0x66, 0x66, 0x66))
        para(doc, "Unit Vocabulary Bank", bold=True, size=12.5, space_before=10)
        vb_table = doc.add_table(rows=0, cols=2)
        vb_table.style = "Table Grid"
        for v in meta.get("vocabulary_bank", []):
            row = vb_table.add_row()
            word_cell_text = v["word"] + (f"  ({v['syllables']})" if v.get("syllables") else "")
            row.cells[0].text = word_cell_text
            row.cells[0].paragraphs[0].runs[0].bold = True
            def_text = v.get("definition", "")
            if v.get("morphology"):
                def_text += f"  —  Word study: {v['morphology']}"
            row.cells[1].text = def_text
        doc.add_page_break()

        for l in lessons:
            heading_bar(doc, f"DAY {l['day']}  ·  {l['title']}", primary_hexc)

            para(doc, "Objectives", bold=True, size=11.5, color=RGBColor.from_string(primary_dk_hexc))
            for o in l.get("objectives", []):
                para(doc, f"☐  {o}", size=10.5, space_after=2)

            para(doc, "Vocabulary", bold=True, size=11.5, color=RGBColor.from_string(primary_dk_hexc), space_before=10)
            for v in l.get("vocab", []):
                word_line = v['word']
                if v.get("syllables"):
                    word_line += f"  ({v['syllables']})"
                para(doc, f"{word_line}: {v.get('definition','')}", size=10.5, space_after=2)
                if v.get("morphology"):
                    para(doc, f"   Word study: {v['morphology']}", italic=True, size=9.5, color=RGBColor(0x66, 0x66, 0x66), space_after=2)
                if v.get("example"):
                    para(doc, f"   e.g. {v['example']}", italic=True, size=10, color=RGBColor(0x66, 0x66, 0x66), space_after=4)

            av = l.get("academic_vocab")
            if av:
                para(doc, "Academic Vocabulary Word of the Day", bold=True, size=11.5, color=RGBColor.from_string(primary_dk_hexc), space_before=10)
                aw_line = av['word']
                if av.get("syllables"):
                    aw_line += f"  ({av['syllables']})"
                para(doc, aw_line, bold=True, size=11, space_after=2)
                para(doc, av.get("definition", ""), size=10.5, space_after=2)
                if av.get("morphology"):
                    para(doc, f"Word study: {av['morphology']}", italic=True, size=9.5, color=RGBColor(0x66, 0x66, 0x66), space_after=2)
                if av.get("example"):
                    para(doc, f"e.g. {av['example']}", italic=True, size=10, color=RGBColor(0x66, 0x66, 0x66), space_after=4)

            para(doc, "Warm-Up", bold=True, size=11.5, color=RGBColor.from_string(primary_dk_hexc), space_before=10)
            para(doc, l.get("warm_up", ""), size=10.5)
            blank_lines(doc, 2)

            para(doc, "Guided Notes", bold=True, size=11.5, color=RGBColor.from_string(primary_dk_hexc), space_before=10)
            for section in l.get("mini_lesson", []):
                para(doc, section.get("heading", ""), bold=True, size=10.5, space_after=2)
                blank_lines(doc, 2)

            act = l.get("activity") or {}
            para(doc, "Activity", bold=True, size=11.5, color=RGBColor.from_string(primary_dk_hexc), space_before=10)
            para(doc, act.get("title", ""), bold=True, size=10.5, space_after=2)
            para(doc, act.get("instructions", ""), size=10.5)

            reading = l.get("reading")
            if reading:
                doc.add_page_break()
                para(doc, f"Reading · {reading.get('type','').replace('_',' ').title()}", bold=True, size=11.5,
                     color=RGBColor.from_string(primary_dk_hexc))
                para(doc, reading.get("title", ""), bold=True, size=12, space_after=6)
                para(doc, reading.get("body", ""), size=10.5)
                if reading.get("questions"):
                    para(doc, "Questions", bold=True, size=11, space_before=8)
                    for q in reading["questions"]:
                        para(doc, f"• {q}", size=10.5, space_after=2)
                        blank_lines(doc, 1)

            app = l.get("application")
            if app:
                para(doc, f"Apply It · {app.get('type','').replace('_',' ').title()}", bold=True, size=11.5,
                     color=RGBColor.from_string(primary_dk_hexc), space_before=12)
                para(doc, app.get("title", ""), bold=True, size=11, space_after=6)
                para(doc, app.get("scenario", ""), size=10.5, space_after=6)
                para(doc, f"Task: {app.get('task','')}", bold=True, size=10.5, space_after=2)
                blank_lines(doc, 3)

            quiz = l.get("quiz") or {}
            if quiz.get("questions"):
                para(doc, "Quiz", bold=True, size=11.5, color=RGBColor.from_string(primary_dk_hexc), space_before=12)
                for i, q in enumerate(quiz["questions"], 1):
                    para(doc, f"{i}. {q['q']}", size=10.5, space_after=2)
                    if q.get("choices"):
                        for c in q["choices"]:
                            para(doc, f"     {c}", size=10.5, space_after=1)
                    else:
                        blank_lines(doc, 2)

            exit_t = l.get("exit_ticket") or {}
            para(doc, "Exit Ticket", bold=True, size=11.5, color=RGBColor(0xFF, 0xFF, 0xFF), space_before=12)
            et_table = doc.add_table(rows=1, cols=1)
            cell = et_table.rows[0].cells[0]
            set_cell_bg(cell, primary_dk_hexc)
            cp = cell.paragraphs[0]
            r = cp.add_run(exit_t.get("prompt", ""))
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r.font.size = Pt(10.5)
            r.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            blank_lines(doc, 3)

            doc.add_page_break()

        # Unit assessment
        ua = meta.get("unit_assessment") or {}
        if ua.get("questions"):
            heading_bar(doc, ua.get("title", f"Unit {meta['number']} Assessment"), primary_dk_hexc)
            for i, q in enumerate(ua["questions"], 1):
                para(doc, f"{i}. {q['q']}", size=10.5, space_after=2)
                if q.get("choices"):
                    for c in q["choices"]:
                        para(doc, f"     {c}", size=10.5, space_after=1)
                else:
                    blank_lines(doc, 2)
            doc.add_page_break()

    doc.save(output_path)
    return output_path
