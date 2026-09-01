"""
Teenager 101 -- Student Workbook Builder (.docx)
Produces a printable workbook: Cornell notes, bilingual reading + questions,
discussion story + questions, vocabulary, and an exit ticket, per lesson.

Run from inside the teen101/ directory: python3 build_workbook.py
"""

import re

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from content import UNITS, SCHOOL_NAME, COURSE_TITLE, RED, RED_DK, GOLD_LT, GRAY, MUTED
from stories import STORIES
from stories_es import STORIES_ES
from readings import READINGS
from readings_es import READINGS_ES

OUT_PATH = "output/Teenager_101_Student_Workbook.docx"


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def strip_html(text):
    return re.sub(r"<[^>]+>", "", text).strip()


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def add_heading(doc, text, size=20, color=RED_DK, space_before=18, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    return p


def add_body(doc, text, size=11, italic=False, color="222222"):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    run.font.color.rgb = rgb(color)
    return p


def add_rule(doc, color=RED):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:color"), color)
    border.append(bottom)
    pPr.append(border)
    p.paragraph_format.space_after = Pt(2)


def title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COURSE_TITLE)
    run.bold = True
    run.font.size = Pt(44)
    run.font.color.rgb = rgb(RED_DK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Student Workbook")
    run.font.size = Pt(22)
    run.font.color.rgb = rgb("444444")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Units 1-6 -- All 40 Lessons")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = rgb(RED)

    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(SCHOOL_NAME)
    run.font.size = Pt(13)
    run.font.color.rgb = rgb(MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Name: _______________________________     Class period: _______")
    run.font.size = Pt(12)

    doc.add_page_break()


def unit_intro(doc, unit):
    add_heading(doc, f"Unit {unit['number']} -- {unit['name']}", size=24)
    add_rule(doc)
    add_body(doc, unit["overview"])
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, label in enumerate(["Lesson", "Title", "Topics"]):
        hdr[i].text = label
        set_cell_shading(hdr[i], RED_DK)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = rgb("FFFFFF")
                r.font.size = Pt(10)
    for lesson in unit["lessons"]:
        row = table.add_row().cells
        row[0].text = lesson["label"]
        row[1].text = lesson["title"]
        row[2].text = lesson["key_topics"]
        for c in row:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_page_break()


def cornell_notes_page(doc, lesson):
    add_heading(doc, f"Lesson {lesson['label']} -- {lesson['title']}", size=18)
    add_rule(doc)
    add_body(doc, "Learning objectives:", size=11, color=MUTED)
    for obj in lesson["objectives"]:
        doc.add_paragraph(obj, style="List Bullet")

    add_heading(doc, "Vocabulary", size=13, color=RED, space_before=14)
    vt = doc.add_table(rows=1, cols=2)
    vt.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = vt.rows[0].cells
    hdr[0].text, hdr[1].text = "Word", "Definition"
    for c in hdr:
        set_cell_shading(c, GRAY)
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for v in lesson["vocabulary"]:
        row = vt.add_row().cells
        row[0].text = v["word"]
        row[1].text = v["definition"]
        for c in row:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)

    add_heading(doc, "Cornell Notes", size=13, color=RED, space_before=16)
    add_body(doc, "Fill in the cue column as you watch the video and discuss as a class. Use the note-taking area for details, then write a 2-3 sentence summary at the bottom.", size=9, italic=True, color=MUTED)

    nt = doc.add_table(rows=1, cols=2)
    nt.alignment = WD_TABLE_ALIGNMENT.CENTER
    nt.columns[0].width = Inches(2.0)
    nt.columns[1].width = Inches(4.5)
    hdr = nt.rows[0].cells
    hdr[0].text, hdr[1].text = "Cues / Questions", "Notes"
    for c in hdr:
        set_cell_shading(c, RED_DK)
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = rgb("FFFFFF")
                r.font.size = Pt(10)
    for q in lesson["cornell_questions"]:
        row = nt.add_row().cells
        row[0].text = q
        row[1].text = ""
        row[0].paragraphs[0].runs[0].font.size = Pt(10)
        for _ in range(3):
            row[1].add_paragraph("_" * 55)
        for p in row[1].paragraphs:
            if p.runs:
                p.runs[0].font.size = Pt(10)
                p.runs[0].font.color.rgb = rgb("CCCCCC")

    add_heading(doc, "Summary", size=11, color=RED, space_before=14)
    for _ in range(4):
        doc.add_paragraph("_" * 90)
    doc.add_page_break()


def reading_page(doc, label, lesson):
    en = READINGS.get(label)
    es = READINGS_ES.get(label)
    if not en:
        return
    add_heading(doc, f"Lesson {label} -- Informational Reading", size=16)
    add_body(doc, f'"{en["title"]}"', italic=True, color=RED)
    add_body(doc, f"Source: {en['source']}", size=9, color=MUTED)
    for para in en["text"]:
        doc.add_paragraph(para)

    add_heading(doc, "Comprehension Questions", size=12, color=RED, space_before=12)
    for i, q in enumerate(en["questions"], 1):
        doc.add_paragraph(f"{i}. {q}")
        for _ in range(2):
            doc.add_paragraph("_" * 90)

    if es:
        doc.add_page_break()
        add_heading(doc, f"Lección {label} -- Lectura Informativa (Español)", size=16)
        add_body(doc, f'"{es["title"]}"', italic=True, color=RED)
        add_body(doc, f"Fuente: {es['source']}", size=9, color=MUTED)
        for para in es["text"]:
            doc.add_paragraph(para)
        add_heading(doc, "Preguntas de Comprensión", size=12, color=RED, space_before=12)
        for i, q in enumerate(es["questions"], 1):
            doc.add_paragraph(f"{i}. {q}")

    p = add_body(doc, f"\U0001F3A7 Audio available: reading_EN_{label}.m4a / reading_ES_{label}.m4a", size=9, color=MUTED)
    doc.add_page_break()


def story_page(doc, label, lesson):
    en = STORIES.get(label)
    if not en:
        return
    add_heading(doc, f"Lesson {label} -- Discussion Story", size=16)
    add_body(doc, f'"{en["title"]}"', italic=True, color=RED)
    add_body(doc, f"Characters: {', '.join(en['characters'])}", size=9, color=MUTED)
    for para in strip_html(en["story"]).split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    add_heading(doc, "Discussion Questions", size=12, color=RED, space_before=12)
    for i, q in enumerate(en["questions"], 1):
        doc.add_paragraph(f"{i}. {q}")
        for _ in range(2):
            doc.add_paragraph("_" * 90)

    add_body(doc, f"\U0001F3A7 Audio available: story_EN_{label}.m4a / story_ES_{label}.m4a", size=9, color=MUTED)

    add_heading(doc, "Exit Ticket", size=12, color=RED, space_before=16)
    doc.add_paragraph(lesson["exit_ticket"])
    for _ in range(3):
        doc.add_paragraph("_" * 90)
    doc.add_page_break()


def build():
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title_page(doc)

    for unit in UNITS:
        unit_intro(doc, unit)
        for lesson in unit["lessons"]:
            label = lesson["label"]
            cornell_notes_page(doc, lesson)
            reading_page(doc, label, lesson)
            story_page(doc, label, lesson)

    import os
    os.makedirs("output", exist_ok=True)
    doc.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    build()
