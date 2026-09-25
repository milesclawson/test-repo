"""
Teenager 101 -- Teacher Slide Deck Builder (.pptx)
Google-Slides-compatible deck: title slide, per-lesson objectives/vocab,
YouTube video slide with watch-for prompt, Cornell notes frame, discussion
story + prompts, and exit ticket, per lesson.

Run from inside the teen101/ directory: python3 build_slides.py
"""

import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from content import UNITS, SCHOOL_NAME, COURSE_TITLE, RED, RED_DK, GOLD, GRAY
from stories import STORIES
from readings import READINGS
from teen101_youtube_links import YOUTUBE

OUT_PATH = "output/Teenager_101_Teacher_Slides.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def strip_html(text):
    return re.sub(r"<[^>]+>", "", text).strip()


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_rect(slide, x, y, w, h, color, hyperlink=None):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    if hyperlink:
        shape.click_action.hyperlink.address = hyperlink
    return shape


def add_text(slide, x, y, w, h, text, size=18, bold=False, italic=False,
             color="222222", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font="Calibri", wrap=True, line_spacing=None, hyperlink=None,
             underline=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
        run.font.name = font
        run.font.underline = underline
        if hyperlink:
            run.hyperlink.address = hyperlink
    return box


def add_bullets(slide, x, y, w, h, items, size=16, color="222222", bold_first=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        p.space_after = Pt(8)
    return box


def header_bar(slide, label, title):
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.0), RED_DK)
    add_text(slide, Inches(0.5), Inches(0.12), Inches(11), Inches(0.4),
              f"TEENAGER 101  |  LESSON {label}", size=14, bold=True, color="FFFFFF")
    add_text(slide, Inches(0.5), Inches(0.45), Inches(11.5), Inches(0.5),
              title, size=26, bold=True, color="FFFFFF")


def footer(slide, text=SCHOOL_NAME):
    add_text(slide, Inches(0.5), Inches(7.15), Inches(8), Inches(0.3),
              text, size=10, color="999999")


# ── Slide builders ──────────────────────────────────────────────────────

def title_slide(prs):
    slide = blank_slide(prs)
    add_bg(slide, RED_DK)
    add_text(slide, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.4),
              COURSE_TITLE, size=60, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.8), Inches(3.9), Inches(11.7), Inches(0.7),
              "Teacher Slide Deck  —  Units 1-6, All 40 Lessons", size=24, color="F5B800", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.5),
              SCHOOL_NAME, size=16, color="FFFFFF", align=PP_ALIGN.CENTER)


def unit_divider(prs, unit):
    slide = blank_slide(prs)
    add_bg(slide, RED)
    add_text(slide, Inches(0.8), Inches(2.9), Inches(11.7), Inches(0.8),
              f"UNIT {unit['number']}", size=28, bold=True, color="F5E9A8", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.8), Inches(3.5), Inches(11.7), Inches(1.0),
              unit["name"], size=48, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.8), Inches(4.7), Inches(9.7), Inches(1.5),
              unit["overview"], size=16, color="FFFFFF", align=PP_ALIGN.CENTER)


def objectives_slide(prs, lesson):
    slide = blank_slide(prs)
    add_bg(slide, "FFFFFF")
    header_bar(slide, lesson["label"], lesson["title"])
    add_text(slide, Inches(0.5), Inches(1.3), Inches(6), Inches(0.4),
              "Learning Objectives", size=18, bold=True, color=RED_DK)
    add_bullets(slide, Inches(0.6), Inches(1.8), Inches(6), Inches(3), lesson["objectives"], size=16)

    add_text(slide, Inches(7.0), Inches(1.3), Inches(5.8), Inches(0.4),
              "Vocabulary", size=18, bold=True, color=RED_DK)
    y = Inches(1.8)
    for v in lesson["vocabulary"]:
        add_text(slide, Inches(7.0), y, Inches(5.8), Inches(0.35), v["word"], size=15, bold=True, color=RED)
        y = Emu(y + Inches(0.38))
        add_text(slide, Inches(7.0), y, Inches(5.8), Inches(0.6), v["definition"], size=12, color="444444")
        y = Emu(y + Inches(0.68))
    add_text(slide, Inches(0.5), Inches(6.4), Inches(9), Inches(0.5),
              f"Key topics: {lesson['key_topics']}", size=13, italic=True, color="666666")
    footer(slide)


def video_slide(prs, label, lesson):
    yt = YOUTUBE.get(label)
    slide = blank_slide(prs)
    add_bg(slide, "FFFFFF")
    if yt and yt.get("alt_watch_for"):
        header_title = f"Videos (2): {yt['title']} + {yt['alt_title']}"
    else:
        header_title = "Video: " + (yt["title"] if yt else "")
    header_bar(slide, label, header_title)
    if not yt:
        footer(slide)
        return
    video_url = yt.get("url", "")
    alt_url = yt.get("alt_url", "")

    add_rect(slide, Inches(0.6), Inches(1.5), Inches(5.6), Inches(3.2), "222222", hyperlink=video_url)
    add_text(slide, Inches(0.6), Inches(2.55), Inches(5.6), Inches(0.8),
              "▶", size=54, color="FFFFFF", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.6), Inches(3.5), Inches(5.6), Inches(0.4),
              f"{yt['title']}", size=14, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.6), Inches(3.85), Inches(5.6), Inches(0.4),
              f"{yt['channel']}  ·  {yt['duration']}", size=12, color="CCCCCC", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.6), Inches(4.85), Inches(5.6), Inches(0.4),
              "Click to search this video on YouTube ↗", size=11, italic=True,
              color="666666", align=PP_ALIGN.CENTER, hyperlink=video_url, underline=True)

    alt_watch_for = yt.get("alt_watch_for", "")

    if alt_watch_for:
        # This lesson needs two videos to cover its full set of objectives -- both are
        # required viewing, not primary + fallback. See the comment in
        # teen101_youtube_links.py for which lessons this applies to and why.
        add_text(slide, Inches(6.5), Inches(1.6), Inches(6.3), Inches(0.35),
                  "Video 1 -- watch for:", size=13, bold=True, color=RED_DK)
        add_text(slide, Inches(6.5), Inches(1.95), Inches(6.3), Inches(1.5),
                  yt["watch_for"], size=12, color="333333")

        add_rect(slide, Inches(6.5), Inches(3.55), Inches(6.3), Inches(2.05), "FFF8E1", hyperlink=alt_url)
        add_text(slide, Inches(6.7), Inches(3.65), Inches(5.9), Inches(0.3),
                  "Video 2 (also required):", size=11, bold=True, color="B8860B")
        add_text(slide, Inches(6.7), Inches(3.95), Inches(5.9), Inches(0.4),
                  f"{yt['alt_title']} — {yt['alt_channel']}", size=12, bold=True, color="444444",
                  hyperlink=alt_url, underline=True)
        add_text(slide, Inches(6.7), Inches(4.35), Inches(5.9), Inches(1.2),
                  alt_watch_for, size=11, color="444444")
    else:
        add_text(slide, Inches(6.5), Inches(1.6), Inches(6.3), Inches(0.4),
                  "Watch for:", size=16, bold=True, color=RED_DK)
        add_text(slide, Inches(6.5), Inches(2.05), Inches(6.3), Inches(2.0),
                  yt["watch_for"], size=15, color="333333")

        add_rect(slide, Inches(6.5), Inches(4.3), Inches(6.3), Inches(1.3), "FFF8E1", hyperlink=alt_url)
        add_text(slide, Inches(6.7), Inches(4.4), Inches(5.9), Inches(0.3),
                  "Backup video (click to search):", size=11, bold=True, color="B8860B")
        add_text(slide, Inches(6.7), Inches(4.7), Inches(5.9), Inches(0.8),
                  f"{yt['alt_title']} — {yt['alt_channel']}", size=12, color="444444",
                  hyperlink=alt_url, underline=True)
    footer(slide)


def cornell_frame_slide(prs, lesson):
    slide = blank_slide(prs)
    add_bg(slide, "FFFFFF")
    header_bar(slide, lesson["label"], "Cornell Notes")
    add_text(slide, Inches(0.5), Inches(1.25), Inches(11.5), Inches(0.4),
              "Project this while students fill in their workbook.", size=13, italic=True, color="666666")

    add_rect(slide, Inches(0.6), Inches(1.8), Inches(4.2), Inches(4.9), "F4F4F4")
    add_text(slide, Inches(0.75), Inches(1.9), Inches(3.9), Inches(0.4), "Cues", size=15, bold=True, color=RED_DK)
    y = Inches(2.4)
    for q in lesson["cornell_questions"]:
        add_text(slide, Inches(0.75), y, Inches(3.9), Inches(0.9), f"• {q}", size=13, color="333333")
        y = Emu(y + Inches(0.95))

    add_rect(slide, Inches(5.0), Inches(1.8), Inches(7.8), Inches(4.9), "FFFFFF")
    add_text(slide, Inches(5.0), Inches(1.8), Inches(7.8), Inches(4.9), "", size=13)
    box = slide.shapes.add_shape(1, Inches(5.0), Inches(1.8), Inches(7.8), Inches(4.9))
    box.fill.background()
    box.line.color.rgb = rgb(GRAY)
    box.line.width = Pt(1.5)
    add_text(slide, Inches(5.15), Inches(1.9), Inches(7.5), Inches(0.4), "Notes", size=15, bold=True, color=RED_DK)
    footer(slide)


def story_slide(prs, label, lesson):
    s = STORIES.get(label)
    if not s:
        return
    slide = blank_slide(prs)
    add_bg(slide, "FFFFFF")
    header_bar(slide, label, f"Discussion Story: “{s['title']}”")
    add_text(slide, Inches(0.5), Inches(1.25), Inches(11.5), Inches(0.35),
              f"Characters: {', '.join(s['characters'])}   |   \U0001F3A7 story_EN_{label}.m4a / story_ES_{label}.m4a",
              size=13, italic=True, color="666666")
    add_text(slide, Inches(0.5), Inches(1.7), Inches(6.3), Inches(0.4),
              "Discussion Questions", size=17, bold=True, color=RED_DK)
    add_bullets(slide, Inches(0.6), Inches(2.15), Inches(6.2), Inches(4.5), s["questions"], size=15)

    reading = READINGS.get(label)
    if reading:
        add_text(slide, Inches(7.1), Inches(1.7), Inches(5.7), Inches(0.4),
                  "Paired reading", size=17, bold=True, color=RED_DK)
        add_text(slide, Inches(7.1), Inches(2.15), Inches(5.7), Inches(0.5),
                  f"“{reading['title']}”", size=14, italic=True, color="333333")
        add_text(slide, Inches(7.1), Inches(2.65), Inches(5.7), Inches(0.4),
                  f"\U0001F3A7 reading_EN_{label}.m4a / reading_ES_{label}.m4a", size=12, color="666666")
    footer(slide)


def exit_ticket_slide(prs, lesson):
    slide = blank_slide(prs)
    add_bg(slide, RED)
    add_text(slide, Inches(0.8), Inches(2.2), Inches(11.7), Inches(0.6),
              f"Exit Ticket — Lesson {lesson['label']}", size=26, bold=True, color="F5E9A8", align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.5), Inches(3.2), Inches(10.3), Inches(2.5),
              lesson["exit_ticket"], size=22, color="FFFFFF", align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title_slide(prs)

    for unit in UNITS:
        unit_divider(prs, unit)
        for lesson in unit["lessons"]:
            label = lesson["label"]
            objectives_slide(prs, lesson)
            video_slide(prs, label, lesson)
            cornell_frame_slide(prs, lesson)
            story_slide(prs, label, lesson)
            exit_ticket_slide(prs, lesson)

    os.makedirs("output", exist_ok=True)
    prs.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    build()
